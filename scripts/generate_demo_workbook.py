"""Generates the fictional demo data for OurFault.

* demo/operations-log-demo.xlsx - an operations log to copy rows from. The rows
  of the demo incident are filled yellow only to make the demo easy to follow;
  OurFault ignores formatting and only uses what the operator pastes.
* demo/demo-paste.txt - the 08:14-10:02 block exactly as Excel puts it on the
  clipboard (tab-separated, CRLF). Used by the "paste example" button and tests.

Usage (requires openpyxl):  python scripts/generate_demo_workbook.py
"""

from datetime import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

DEMO_DIR = Path(__file__).resolve().parent.parent / "demo"
OUTPUT = DEMO_DIR / "operations-log-demo.xlsx"
PASTE_OUTPUT = DEMO_DIR / "demo-paste.txt"

YELLOW = PatternFill("solid", fgColor="FFFF00")
GREEN = PatternFill("solid", fgColor="E2EFDA")
HEADER = PatternFill("solid", fgColor="D9D9D9")

# (time, from, to, description, event type, fill)
ROWS = [
    ((7, 0), "מוקד מבצעים", "כלל העמדות", "פתיחת משמרת בוקר. כל העמדות מאוישות.", "שגרה", None),
    ((7, 12), "עמדה 2", "מוקד מבצעים", "בדיקת קשר תקופתית הושלמה בהצלחה.", "שגרה", None),
    ((7, 25), "מוקד מבצעים", "תורן טכני", "תזכורת: עבודות תחזוקה מתוכננות בשרת הגיבוי בשעה 10:00.", "תיאום", None),
    ((7, 40), "עמדה 4", "מוקד מבצעים", "עדכון מזג אוויר: ראות טובה, ללא התרעות.", "שגרה", None),
    ((7, 58), "עמדה 1", "מוקד מבצעים", "התקבלה בקשה לתיאום הגעת צוות אחזקה בשעה 09:30.", "תיאום", None),
    ((8, 14), "עמדה 3", "מוקד מבצעים", "זוהתה האטה בקצב עדכון הנתונים במערכת אלפא (עיכוב של כ-40 שניות).", "תקלה", YELLOW),
    ((8, 16), "מוקד מבצעים", "תורן טכני", "דווח לתורן הטכני על האטה בעדכוני מערכת אלפא. התבקשה בדיקה.", "תקלה", YELLOW),
    ((8, 22), "עמדה 2", "מוקד מבצעים", "בדיקת קשר תקופתית הושלמה בהצלחה.", "שגרה", None),
    ((8, 31), "תורן טכני", "מוקד מבצעים", "אותר עומס חריג על שרת התקשורת. בוצע מעבר לשרת החלופי.", "תקלה", YELLOW),
    ((8, 45), "עמדה 4", "מוקד מבצעים", "סיום סבב סיור שגרתי. ללא חריגים.", "שגרה", None),
    ((8, 52), "עמדה 3", "מוקד מבצעים", "קצב עדכון הנתונים במערכת אלפא חזר לשגרה.", "עדכון", YELLOW),
    ((9, 5), "מוקד מבצעים", "קצין תורן", "הועבר דיווח ראשוני על התקלה במערכת אלפא ומשך ההשפעה (כ-38 דקות).", "דיווח", YELLOW),
    ((9, 30), "צוות אחזקה", "מוקד מבצעים", "צוות האחזקה הגיע לאתר והחל בעבודה.", "תיאום", GREEN),
    ((9, 48), "עמדה 1", "מוקד מבצעים", "התקבלה הודעה על תרגיל מתוכנן ביום חמישי.", "מידע", None),
    ((10, 2), "תורן טכני", "מוקד מבצעים", "בדיקות מקדימות לתקלה בשרת התקשורת הושלמו. הממצאים תועדו בנפרד.", "עדכון", YELLOW),
    ((10, 15), "מוקד מבצעים", "כלל העמדות", "תחזוקת שרת הגיבוי החלה כמתוכנן.", "תיאום", None),
    ((10, 40), "עמדה 2", "מוקד מבצעים", "בדיקת קשר תקופתית הושלמה בהצלחה.", "שגרה", None),
    ((11, 5), "מוקד מבצעים", "תורן טכני", "תחזוקת שרת הגיבוי הסתיימה. המערכות פועלות כסדרן.", "עדכון", None),
    ((11, 20), "עמדה 4", "מוקד מבצעים", "יציאה לסבב סיור שני.", "שגרה", None),
    ((11, 45), "מוקד מבצעים", "כלל העמדות", "תדריך אמצע משמרת הועבר לכלל העמדות.", "שגרה", None),
]


def main() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "יומן מבצעים"
    sheet.sheet_view.rightToLeft = True

    sheet.append(["יומן מבצעים – משמרת בוקר – 23/09/2026 (נתונים לדוגמה בלבד)"])
    sheet["A1"].font = Font(bold=True, size=13)

    sheet.append(["Time", "From", "To", "Description", "Event type"])
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = HEADER

    for (hour, minute), sender, recipient, description, event_type, fill in ROWS:
        sheet.append([time(hour, minute), sender, recipient, description, event_type])
        row = sheet[sheet.max_row]
        row[0].number_format = "hh:mm"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill is not None:
                cell.fill = fill

    for column, width in zip("ABCDE", (9, 16, 16, 70, 12)):
        sheet.column_dimensions[column].width = width

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(f"wrote {OUTPUT}")

    # Excel copies the displayed cell text, tab-separated, one CRLF-terminated line per row.
    block = [row for row in ROWS if (8, 14) <= row[0] <= (10, 2)]
    lines = [f"{hour:02}:{minute:02}\t{sender}\t{recipient}\t{description}\t{event_type}"
             for (hour, minute), sender, recipient, description, event_type, _ in block]
    PASTE_OUTPUT.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))
    print(f"wrote {PASTE_OUTPUT}")


if __name__ == "__main__":
    main()
